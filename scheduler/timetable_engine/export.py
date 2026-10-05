"""Export: timetable.json (for the AMS, BS/MS switchable) + Excel workbook with all views."""
import json, collections, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from engine import CFG, ROOMS

HEAD = PatternFill('solid', fgColor='1F3864')
SUB = PatternFill('solid', fgColor='D9E2F3')
BAD = PatternFill('solid', fgColor='F8CBAD')
WHITE = Font(color='FFFFFF', bold=True)
THIN = Side(style='thin', color='A6A6A6')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)


GSIZE = {}


def entries(tt, offerings):
    out = []
    for si, st in sorted(tt.start.items(), key=lambda kv: (kv[1], kv[0])):
        ss = tt.sessions[si]
        sec = tt.secmap[ss['sec']]
        off = offerings[sec['off']]
        r = ROOMS[tt.room[si]]
        d, k = divmod(st, tt.NS)
        groups = collections.Counter(tt.students[i]['cohort'] for i in sec['final'])
        gcnt = collections.Counter(tt.students[i]['group'] for i in sec['final'])
        gsize = GSIZE.setdefault(id(tt), collections.Counter(s['group'] for s in tt.students))
        classes = [g for g, n in gcnt.items() if n >= max(3, 0.3 * gsize[g])]
        out.append({
            'timetable': tt.name, 'course_code': off['code'], 'course_title': off['title'],
            'section': sec['label'], 'section_uid': f"{off['code']}-{sec['label']}",
            'teacher': sec['teacher'] or 'TBA', 'theory_credit_hours': sec['theory_cr'],
            'tag': sec['tag'], 'session_no': ss['j'] + 1, 'duration_slots': ss['dur'],
            'day': tt.days[d], 'slot_index': k, 'time': tt.slots[k] if ss['dur'] == 1 else
            f"{tt.slots[k].split('-')[0]}-{tt.slots[k + ss['dur'] - 1].split('-')[1]}",
            'room': r['name'], 'block': r['block'], 'room_capacity': r['cap'],
            'students': len(sec['final']), 'type': 'theory',
            'cohorts': [f'{c} ({n})' for c, n in groups.most_common()],
            'classes': classes,
            'student_ids': sorted(tt.students[i]['id'] for i in sec['final']),
        })
    return out


def metrics(tt):
    m = {}
    m['sections'] = len(tt.secs)
    m['sessions'] = len(tt.sessions)
    m['room_slots_available'] = len(ROOMS) * tt.T
    m['room_utilisation_pct'] = round(100 * len(tt.sessions) / (len(ROOMS) * tt.T), 1)
    rep = tt.report
    for k in ('room_double_booked', 'teacher_clash', 'student_clash', 'same_day_repeat'):
        m[k] = len(rep.get(k, []))
    m['sessions_over_room_capacity'] = len(rep['over_capacity'])
    m['students_seated_beyond_capacity'] = sum(n - c for _, n, c in rep['over_capacity'])
    two = [s for s in tt.secs if len(s['sess']) == 2]
    gaps = collections.Counter()
    same_time = 0
    for s in two:
        a, b = sorted(tt.start[si] for si in s['sess'])
        gaps[b // tt.NS - a // tt.NS] += 1
        same_time += (a % tt.NS == b % tt.NS)
    m['2x_sections_day_gap'] = dict(sorted(gaps.items()))
    m['2x_sections_same_time_both_days'] = same_time
    # per student daily load
    per = collections.defaultdict(collections.Counter)
    for s in tt.secs:
        for si in s['sess']:
            for i in s['final']:
                per[i][tt.start[si] // tt.NS] += 1
    loads = collections.Counter(max(c.values()) for c in per.values())
    m['students_by_max_classes_per_day'] = dict(sorted(loads.items()))
    m['students_total'] = len(per)
    busy = collections.Counter()
    for si, st in tt.start.items():
        for t in range(st, st + tt.sessions[si]['dur']):
            busy[t] += 1
    m['rooms_busy_per_slot'] = {f'{tt.days[t // tt.NS]} {tt.slots[t % tt.NS]}': busy[t]
                                for t in range(tt.T)}
    m['solver_status'] = tt.status
    return m


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD, WHITE, CENTER, BOX


def grid_sheet(wb, title, tt, rows_key, rows_label, cell_text, rows_order=None):
    """generic weekly grid: one row per entity, columns = day x slot."""
    ws = wb.create_sheet(title[:31])
    ws.cell(row=1, column=1, value=rows_label)
    col = 2
    for d in tt.days:
        for s in tt.slots:
            ws.cell(row=1, column=col, value=f'{d}\n{s}')
            col += 1
    style_header(ws, 1, col - 1)
    cells = collections.defaultdict(list)
    for si, st in tt.start.items():
        for key in rows_key(si):
            for t in range(st, st + tt.sessions[si]['dur']):
                cells[key, t].append(cell_text(si))
    keys = rows_order or sorted({k for k, _ in cells})
    for r, key in enumerate(keys, start=2):
        ws.cell(row=r, column=1, value=key).font = Font(bold=True)
        ws.cell(row=r, column=1).fill = SUB
        ws.cell(row=r, column=1).border = BOX
        for t in range(tt.T):
            txt = '\n'.join(cells.get((key, t), []))
            c = ws.cell(row=r, column=2 + t, value=txt or None)
            c.alignment, c.border = WRAP, BOX
            if len(cells.get((key, t), [])) > 1:
                c.fill = BAD
    ws.column_dimensions['A'].width = 34
    for c in range(2, tt.T + 2):
        ws.column_dimensions[get_column_letter(c)].width = 22
    ws.freeze_panes = 'B2'
    ws.row_dimensions[1].height = 32
    return ws


def export(students, offerings, sections, results, log):
    payload = {'generated_at': datetime.datetime.now().isoformat(timespec='seconds'),
               'institution': 'Abasyn University Islamabad Campus', 'term': 'Fall 2026',
               'rooms': ROOMS, 'timetables': {}}
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    r = 1
    ws.cell(row=r, column=1, value='AUIC Theory Timetable - Fall 2026 (generated)').font = Font(bold=True, size=14)
    r += 2
    for name, tt in results.items():
        if not getattr(tt, 'start', None):
            ws.cell(row=r, column=1, value=f'{name}: NO SOLUTION ({tt.status})')
            r += 2
            continue
        E = entries(tt, offerings)
        M = metrics(tt)
        payload['timetables'][name] = {'days': tt.days, 'slots': tt.slots, 'metrics': M,
                                       'entries': E}
        ws.cell(row=r, column=1, value=f'{name} timetable').font = Font(bold=True, size=12)
        r += 1
        for k, v in M.items():
            if k == 'rooms_busy_per_slot':
                continue
            ws.cell(row=r, column=1, value=k)
            ws.cell(row=r, column=2, value=json.dumps(v) if isinstance(v, dict) else v)
            r += 1
        r += 1

        rn = lambda si: ROOMS[tt.room[si]]['name']
        def txt(si):
            s = tt.secmap[tt.sessions[si]['sec']]
            o = offerings[s['off']]
            return f"{o['code']}-{s['label']} {s['tag']}".strip() + f"\n{o['title'][:28]}\n{s['teacher'] or 'TBA'} ({len(s['final'])})"
        def txt_room(si):
            s = tt.secmap[tt.sessions[si]['sec']]
            o = offerings[s['off']]
            return f"{o['code']}-{s['label']} {s['tag']}".strip() + f" @ {rn(si)}"

        grid_sheet(wb, f'{name} Room-wise', tt, lambda si: [rn(si)], 'Room (capacity)',
                   txt, rows_order=[x['name'] for x in ROOMS])
        # rename room rows with capacity
        wsr = wb[f'{name} Room-wise'[:31]]
        for i, x in enumerate(ROOMS, start=2):
            wsr.cell(row=i, column=1, value=f"{x['name']} ({x['cap']})")

        # class-wise: group -> sections where the group is a real part
        def groups_of(si):
            s = tt.secmap[tt.sessions[si]['sec']]
            cnt = collections.Counter(tt.students[i]['group'] for i in s['final'])
            gsize = collections.Counter(st['group'] for st in tt.students)
            return [g for g, n in cnt.items() if n >= max(3, 0.3 * gsize[g])]
        grid_sheet(wb, f'{name} Class-wise', tt, groups_of, 'Program | Intake | Group',
                   txt_room)
        grid_sheet(wb, f'{name} Teacher-wise', tt,
                   lambda si: [tt.secmap[tt.sessions[si]['sec']]['teacher'] or 'TBA (no teacher)'],
                   'Teacher', txt_room)
        # flat list
        wl = wb.create_sheet(f'{name} Session List')
        cols = ['day', 'time', 'room', 'room_capacity', 'course_code', 'section', 'course_title',
                'tag', 'teacher', 'students', 'session_no', 'cohorts']
        for c, h in enumerate(cols, 1):
            wl.cell(row=1, column=c, value=h)
        style_header(wl, 1, len(cols))
        for i, e in enumerate(E, start=2):
            for c, h in enumerate(cols, 1):
                v = e[h]
                wl.cell(row=i, column=c, value=', '.join(v) if isinstance(v, list) else v)
                if h == 'students' and e['students'] > e['room_capacity']:
                    wl.cell(row=i, column=c).fill = BAD
        for c, w in enumerate([6, 12, 20, 8, 10, 7, 40, 8, 28, 9, 8, 60], 1):
            wl.column_dimensions[get_column_letter(c)].width = w
        wl.freeze_panes = 'A2'
        wl.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(E) + 1}'

    # student -> sections
    wsx = wb.create_sheet('Student Sections')
    hdr = ['Student ID', 'Name', 'Program', 'Intake', 'Timetable', 'Course', 'Section', 'Teacher']
    for c, h in enumerate(hdr, 1):
        wsx.cell(row=1, column=c, value=h)
    style_header(wsx, 1, len(hdr))
    rr = 2
    for name, tt in results.items():
        if not getattr(tt, 'start', None):
            continue
        for s in tt.secs:
            o = offerings[s['off']]
            for i in s['final']:
                st = students[i]
                for c, v in enumerate([st['id'], st['name'], st['program'], st['batch'], name,
                                       f"{o['code']} {o['title']}", s['label'], s['teacher'] or 'TBA'], 1):
                    wsx.cell(row=rr, column=c, value=v)
                rr += 1
    for c, w in enumerate([10, 26, 30, 12, 9, 50, 8, 28], 1):
        wsx.column_dimensions[get_column_letter(c)].width = w
    wsx.auto_filter.ref = f'A1:H{rr - 1}'
    wsx.freeze_panes = 'A2'

    # excluded / merged + capacity overflow
    wse = wb.create_sheet('Excluded & Notes')
    r = 1
    for k, items in log.items():
        wse.cell(row=r, column=1, value=f'{k} ({len(items)})').font = Font(bold=True)
        r += 1
        for it in items:
            wse.cell(row=r, column=1, value=it)
            r += 1
        r += 1
    for name, tt in results.items():
        if not getattr(tt, 'start', None):
            continue
        wse.cell(row=r, column=1, value=f'{name}: sessions larger than their room').font = Font(bold=True)
        r += 1
        for si, n, cap in sorted(tt.report['over_capacity'], key=lambda x: x[2] - x[1]):
            s = tt.secmap[tt.sessions[si]['sec']]
            o = offerings[s['off']]
            wse.cell(row=r, column=1, value=f"{o['code']}-{s['label']} {o['title']}: {n} students in "
                                             f"{ROOMS[tt.room[si]]['name']} (cap {cap})")
            r += 1
        r += 1
    wse.column_dimensions['A'].width = 110

    wb.save('AUIC_Theory_Timetable_Fall2026.xlsx')
    json.dump(payload, open('timetable.json', 'w'), indent=1)
    print('exported')
